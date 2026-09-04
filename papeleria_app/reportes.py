"""
reportes.py
-----------
Generación de reportes de ventas en Excel (openpyxl), a partir de los datos
reales en SQLite (SQLAlchemy). Excel es SOLO el formato de salida -- SQLite
sigue siendo la única fuente de datos.

Genera un .xlsx con 3 hojas: Detalle, Resumen, Resumen por producto.
Además guarda automáticamente una copia en la carpeta `reportes/` (se crea
sola si no existe).
"""

import io
from datetime import date, datetime, time
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import Session

from database import DB_PATH
from models import DetalleVenta, Venta

CARPETA_REPORTES = DB_PATH.parent / "reportes"

RELLENO_ENCABEZADO = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FUENTE_ENCABEZADO = Font(color="FFFFFF", bold=True)
FORMATO_MONEDA = '"S/" #,##0.00'
FORMATO_FECHA = "DD/MM/YYYY"


def _descripcion(nombre: str, marca: Optional[str]) -> str:
    """Nombre + marca, sin inventar nada que no esté en la base de datos."""
    return f"{nombre} - {marca}" if marca else nombre


def _obtener_ventas_del_periodo(db: Session, desde: date, hasta: date):
    inicio = datetime.combine(desde, time.min)
    fin = datetime.combine(hasta, time.max)
    return (
        db.query(Venta)
        .filter(Venta.fecha >= inicio, Venta.fecha <= fin)
        .order_by(Venta.fecha)
        .all()
    )


def _autoajustar_columnas(ws):
    for columna in ws.columns:
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = min(largo + 3, 45)


def generar_reporte(db: Session, desde: date, hasta: date) -> tuple[Optional[bytes], Optional[str], dict]:
    """
    Devuelve (bytes_del_excel, nombre_archivo, resumen) o (None, None, {})
    si no hay ventas en el periodo.
    """
    ventas = _obtener_ventas_del_periodo(db, desde, hasta)
    if not ventas:
        return None, None, {}

    # --- Cálculos para el resumen (a nivel de venta, no de línea, para no
    #     contar el total de una venta más de una vez por tener varios
    #     productos) ---
    total_efectivo = sum(float(v.total) for v in ventas if v.metodo_pago == "EFECTIVO")
    total_yape = sum(float(v.total) for v in ventas if v.metodo_pago == "YAPE")
    total_general = total_efectivo + total_yape
    cantidad_ventas = len(ventas)
    unidades_vendidas = sum(d.cantidad for v in ventas for d in v.detalles)

    # --- Agrupación por producto ---
    por_producto: dict[tuple, dict] = {}
    for v in ventas:
        for d in v.detalles:
            clave = (d.producto.nombre, d.producto.marca)
            acumulado = por_producto.setdefault(clave, {"unidades": 0, "total": 0.0})
            acumulado["unidades"] += d.cantidad
            acumulado["total"] += float(d.subtotal)

    wb = Workbook()

    # ================= Hoja "Detalle" =================
    ws = wb.active
    ws.title = "Detalle"
    encabezados = ["DESCRIPCION", "UNIDAD", "FECHA", "HORA", "MONTO", "METODO_PAGO"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.fill = RELLENO_ENCABEZADO
        celda.font = FUENTE_ENCABEZADO
        celda.alignment = Alignment(horizontal="center")

    fila_num = 2
    for v in ventas:
        for d in v.detalles:
            ws.append([
                _descripcion(d.producto.nombre, d.producto.marca),
                d.cantidad,
                v.fecha.date(),
                v.fecha.strftime("%H:%M"),
                float(d.subtotal),
                v.metodo_pago.value if hasattr(v.metodo_pago, "value") else v.metodo_pago,
            ])
            ws.cell(row=fila_num, column=3).number_format = FORMATO_FECHA
            ws.cell(row=fila_num, column=5).number_format = FORMATO_MONEDA
            fila_num += 1

    ultima_fila = fila_num - 1
    if ultima_fila >= 2:
        tabla = Table(displayName="TablaDetalle", ref=f"A1:F{ultima_fila}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
        )
        ws.add_table(tabla)
    ws.freeze_panes = "A2"
    _autoajustar_columnas(ws)

    # Fila de totales al final
    ws.append(["TOTAL", unidades_vendidas, "", "", total_general, ""])
    for col in (1, 5):
        ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=5).number_format = FORMATO_MONEDA

    # ================= Hoja "Resumen" =================
    ws2 = wb.create_sheet("Resumen")
    filas_resumen = [
        ("Periodo", f"{desde.strftime('%d/%m/%Y')} - {hasta.strftime('%d/%m/%Y')}"),
        ("Cantidad de ventas", cantidad_ventas),
        ("Unidades vendidas", unidades_vendidas),
        ("Total vendido", total_general),
        ("Total efectivo", total_efectivo),
        ("Total Yape", total_yape),
    ]
    for etiqueta, valor in filas_resumen:
        ws2.append([etiqueta, valor])
    for fila in (4, 5, 6):
        ws2.cell(row=fila, column=2).number_format = FORMATO_MONEDA
    for celda in ws2["A"]:
        celda.font = Font(bold=True)
    _autoajustar_columnas(ws2)

    # ================= Hoja "Resumen por producto" =================
    ws3 = wb.create_sheet("Resumen por producto")
    ws3.append(["PRODUCTO", "MARCA", "UNIDADES VENDIDAS", "TOTAL VENDIDO"])
    for celda in ws3[1]:
        celda.fill = RELLENO_ENCABEZADO
        celda.font = FUENTE_ENCABEZADO
    for (nombre, marca), datos in sorted(por_producto.items(), key=lambda x: -x[1]["total"]):
        ws3.append([nombre, marca or "", datos["unidades"], datos["total"]])
        ws3.cell(row=ws3.max_row, column=4).number_format = FORMATO_MONEDA
    if len(por_producto) > 0:
        tabla3 = Table(displayName="TablaProductos", ref=f"A1:D{len(por_producto) + 1}")
        tabla3.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws3.add_table(tabla3)
    ws3.freeze_panes = "A2"
    _autoajustar_columnas(ws3)

    # --- Nombre de archivo ---
    if desde == hasta:
        nombre_archivo = f"Reporte_Ventas_{desde.strftime('%Y-%m-%d')}.xlsx"
    else:
        nombre_archivo = f"Reporte_Ventas_{desde.strftime('%d-%m-%Y')}_a_{hasta.strftime('%d-%m-%Y')}.xlsx"

    buffer = io.BytesIO()
    wb.save(buffer)
    contenido = buffer.getvalue()

    # --- Guardar copia automática en reportes/ ---
    CARPETA_REPORTES.mkdir(parents=True, exist_ok=True)
    (CARPETA_REPORTES / nombre_archivo).write_bytes(contenido)

    resumen = {
        "cantidad_ventas": cantidad_ventas,
        "unidades_vendidas": unidades_vendidas,
        "total_vendido": round(total_general, 2),
        "total_efectivo": round(total_efectivo, 2),
        "total_yape": round(total_yape, 2),
    }
    return contenido, nombre_archivo, resumen
